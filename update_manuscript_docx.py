import os
import shutil
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor

def update_manuscript():
    backup_path = r"C:\Users\James\Downloads\SWU 003 (Final Project) System Design & Governance Document.backup.docx"
    src_path = r"C:\Users\James\Downloads\SWU 003 (Final Project) System Design & Governance Document.docx"
    
    if os.path.exists(backup_path):
        shutil.copyfile(backup_path, src_path)
        print("Restored from clean backup.")

    doc = Document(src_path)

    def format_run(run, bold=False, italic=False, size=11):
        run.font.name = 'Arial'
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.italic = italic
        run.font.underline = False
        run.font.color.rgb = RGBColor(0, 0, 0)
    
    print("1. Updating Section 3.5 Database Metadata...")
    p84 = doc.paragraphs[84]
    p84.text = ""
    r1 = p84.add_run(
        "The system utilizes a relational database management system hosted on TiDB Cloud (Serverless MySQL 8.0 compatible engine) "
        "operating with utf8mb4 character encoding and utf8mb4_general_ci collation. The complete schema comprises twenty-seven (27) "
        "normalized tables structured across user role-based access control, donation lifecycles, relief inventory tracking, "
        "kit repacking operations, calamity distribution logistics, volunteer shift scheduling, and cryptographically verifiable digital certificates.\n\n"
        "Database Metadata and Data Dictionary:\n"
        "Document File: database/DATABASE_METADATA.md (located in the project database folder)\n"
        "GitHub Repository Path: github.com/jamis-ops/Donation-Management-System/blob/main/database/DATABASE_METADATA.md\n\n"
        "Master SQL Schema:\n"
        "Database File: database/donation_system.sql\n"
        "GitHub Repository Path: github.com/jamis-ops/Donation-Management-System/blob/main/database/donation_system.sql"
    )
    format_run(r1, bold=False, size=11)
    
    print("2. Updating Section 3.6 Environments (Tables 6, 7, 8)...")
    # Table 6: For the System
    t6 = doc.tables[6]
    system_links = [
        "https://donation-system-frontend.vercel.app\n(Primary Production SPA on Vercel Global Edge CDN)",
        "https://donation-system-uat.vercel.app\n(UAT Staging Environment for Rise Above coordinators)",
        "http://localhost:5173\n(Local Vite and React Development Server)"
    ]
    for c_idx, text in enumerate(system_links):
        cell = t6.rows[1].cells[c_idx]
        cell.text = ""
        p = cell.paragraphs[0]
        r = p.add_run(text)
        format_run(r, bold=False, size=10)

    # Table 7: For the API
    t7 = doc.tables[7]
    api_links = [
        "https://donation-system-api.onrender.com/api\n(Dockerized PHP 8.2 on Render Cloud, Singapore Region)",
        "https://donation-system-api-uat.onrender.com/api\n(UAT Staging API Instance on Render Cloud)",
        "http://localhost:8000/api\n(Local PHP Built-in Server CLI)"
    ]
    for c_idx, text in enumerate(api_links):
        cell = t7.rows[1].cells[c_idx]
        cell.text = ""
        p = cell.paragraphs[0]
        r = p.add_run(text)
        format_run(r, bold=False, size=10)

    # Table 8: For the Database
    t8 = doc.tables[8]
    db_links = [
        "gateway01.ap-southeast-1.prod.aws.tidbcloud.com:4000\nDatabase: donation_system\n(TLS/SSL Encrypted Serverless Cluster)",
        "gateway01.ap-southeast-1.prod.aws.tidbcloud.com:4000\nDatabase: donation_system_uat\n(Isolated UAT Staging Branch)",
        "127.0.0.1:3306 or local_test.sqlite\nDatabase: donation_system_test\n(Local SQLite or MySQL Dev Instance)"
    ]
    for c_idx, text in enumerate(db_links):
        cell = t8.rows[1].cells[c_idx]
        cell.text = ""
        p = cell.paragraphs[0]
        r = p.add_run(text)
        format_run(r, bold=False, size=10)

    print("3. Updating Section 3.7 Service Accounts (Table 9)...")
    # Table 9: For the System
    t9 = doc.tables[9]
    svc_emails = [
        "Primary Super Admin:\nadmin@riseabovefoundation.org\n\nNotification Mailer (Brevo):\nnoreply@riseabovefoundation.org\n\nCloud DevOps Infrastructure:\ndevops@riseabovefoundation.org",
        "UAT Lead Administrator:\nuat-admin@riseabovefoundation.org\n\nUAT Notification Mailer:\nuat-mailer@riseabovefoundation.org\n\nUAT QA Coordinator:\nuat-tester@riseabovefoundation.org",
        "Local Development Superuser:\ntest-admin@localhost.dev\n\nMock Notification Mailer:\nmailhog@localhost.dev\n\nCapstone Lead Developer:\njamesisidro0808@gmail.com"
    ]
    for c_idx, text in enumerate(svc_emails):
        cell = t9.rows[1].cells[c_idx]
        cell.text = ""
        p = cell.paragraphs[0]
        r = p.add_run(text)
        format_run(r, bold=False, size=10)

    print("4. Updating Sections 4.2.2 through 4.2.7 in Governance...")
    
    gov_content = [
        # 4.2.2 Code Naming Convention
        ("heading", "4.2.2. Code Naming Convention"),
        ("body", "The establishment of uniform code naming conventions ensures maintainability, seamless readability, and cross-stack consistency across a heterogeneous architecture combining a modern React Single Page Application (SPA) frontend, a PHP 8.2 RESTful backend API, and a relational MySQL/TiDB database. Consistent naming prevents defects, eases collaborative development among team members, and lowers the cognitive load when inspecting distributed system components."),
        ("bullet_bold", "PascalCase: ", "Exclusively applied to React UI Components (e.g., StatsSection.jsx, PartnersSection.jsx, AdminLayout.jsx), React Context Providers (AuthContext.jsx, ThemeContext.jsx), and PHP Object-Oriented Classes or Exceptions."),
        ("bullet_bold", "camelCase: ", "Applied to JavaScript variables, React custom hooks (e.g., useApiList, usePagination, useAuth), component props, JavaScript helper functions (e.g., fetchDonationStats(), handleVerify()), and JSON API request or response payload attributes."),
        ("bullet_bold", "snake_case: ", "Applied to all database tables and column identifiers (e.g., inventory_items, tracking_code, proof_file, created_at) and backend PHP variables, database query parameters, and procedural helper scripts (e.g., $db_host, get_db_connection())."),
        ("bullet_bold", "UPPER_SNAKE_CASE: ", "Strictly reserved for global constants, configuration definitions, and environment variables (e.g., DB_HOST, DB_PASSWORD, JWT_SECRET, VITE_API_URL, MAX_UPLOAD_SIZE_MB)."),
        ("bullet_bold", "kebab-case: ", "Applied to CSS class names (e.g., .stat-card, .btn-primary, .modal-backdrop), URL web route paths (e.g., /admin/inventory-management, /portals/donor/certificates), and REST API file and resource endpoints."),
        ("body", "Prefixes and Suffixes Requirements:"),
        ("bullet_bold", "Booleans: ", "Must begin with affirmative verbs such as is, has, can, or should (e.g., isLoading, isVerified, hasPermission, canDistribute)."),
        ("bullet_bold", "React Custom Hooks: ", "Must invariably start with the use prefix in compliance with React specifications (e.g., useApiList, usePagination, useAuth)."),
        ("bullet_bold", "Event Handlers: ", "Must use handle for function definitions (e.g., handleDonationSubmit) and on for passed callback props (e.g., onSubmit, onStatusChange)."),
        ("bullet_bold", "Foreign Keys: ", "Must always be suffixed with _id matching the referenced singular entity (e.g., user_id, donor_id, role_id)."),
        ("body", "Rules for Abbreviations: Cryptic abbreviations are strictly prohibited. Standard industry abbreviations are permitted: id, db, api, url, cv, auth, lgu. Business domain entities must use full descriptive terms (e.g., use beneficiary instead of ben, and repacking_job instead of rpk_j)."),
        ("body", "Exceptions: Environment variables and constants are written in ALL_CAPS. Embedded acronyms in camelCase maintain consistent casing (e.g., parseCsvFile, fetchApiData)."),

        # 4.2.3 Coding Standards
        ("heading", "4.2.3. Coding Standards"),
        ("body", "Overall Coding Philosophy: The development team adheres to core software engineering methodologies: 1) DRY (Don't Repeat Yourself) — common database routines, authentication filters, and UI notification modals are centralized into reusable modules; 2) SOLID Principles — especially Single Responsibility for API endpoints and Open/Closed for relief status workflows; 3) KISS (Keep It Simple, Stupid) — avoiding speculative over-engineering in favor of robust, transparent logic; and 4) Defensive Programming — rigorous server-side input sanitization, strict parameter type-casting, and prepared PDO SQL statements to block SQL injection and Cross-Site Scripting (XSS)."),
        ("bullet_bold", "Frontend Style Guide: ", "Airbnb JavaScript Style Guide and Official React.dev best practices."),
        ("bullet_bold", "Backend Style Guide: ", "PHP-FIG PSR-12 (Extended Coding Style Guide) and PSR-4 (Autoloading Standards)."),
        ("bullet_bold", "Database Style Guide: ", "SQL Style Guide (ISO/IEC 9075) with capitalized standard SQL keywords and lowercase table/column identifiers."),
        ("body", "Formatting Rules: JavaScript, JSX, and CSS use 2-space soft indentation; PHP backend code uses 4-space indentation. Maximum line length is capped at 100 characters. One True Brace Style (1TBS) is enforced across JavaScript; PSR-12 brace placement is followed in PHP (classes and methods open on next line; control structures open on same line). Semicolons are mandatory."),
        ("body", "Commenting and Documentation Expectations: All exported functions, hooks, and endpoints must be documented using structured JSDoc (Frontend) or PHPDoc (Backend) indicating parameter data types, expected returns, and exceptions thrown. Inline comments must document architectural intent (why) rather than self-evident code statements (what)."),
        ("body", "Error Handling Conventions: All API responses conform to a unified JSON response schema containing success (boolean), data (payload or null), message (descriptive string), and error (standardized error code or null). Explicit HTTP status codes must be returned (200, 201, 400, 401, 403, 404, 422, 500). Internal database errors are logged securely server-side without leaking raw credentials or stack traces to client browsers."),
        ("body", "Tools Used to Enforce Standards: ESLint and Prettier for frontend formatting; PHP_CodeSniffer (phpcs) for backend PSR-12 validation; Git pre-commit hooks configured via Husky to block non-compliant commits."),

        # 4.2.4 Environment Variables Standards
        ("heading", "4.2.4. Environment Variables Standards"),
        ("body", "Purpose: In strict conformance with the Twelve-Factor App methodology (Config III), all environment-specific configurations, cryptographic secrets, and third-party credentials are completely decoupled from application source code, preventing accidental credential leaks and allowing identical builds to run across Development, UAT, and Production."),
        ("body", "Naming Conventions and Categories: All environment variables use strict UPPER_SNAKE_CASE prefixed by domain:"),
        ("bullet_bold", "Database: ", "DB_HOST, DB_PORT, DB_NAME, DB_USER, DB_PASSWORD, DB_SSL."),
        ("bullet_bold", "Integrations: ", "BREVO_API_KEY, GEMINI_API_KEY, MAIL_FROM_EMAIL, MAIL_FROM_NAME."),
        ("bullet_bold", "Frontend (Vite): ", "VITE_API_URL, VITE_APP_ENV, VITE_APP_VERSION."),
        ("bullet_bold", "CORS and Security: ", "FRONTEND_URL, SESSION_LIFETIME, APP_SECRET."),
        ("body", "Storage and Locations: Locally, environment variables are loaded from an uncommitted .env file. In cloud environments, secrets are managed securely through the Render Environment Dashboard (Backend) and Vercel Project Settings (Frontend). A sanitized template file (.env.example) is committed to version control with descriptive placeholders."),
        ("body", "Hardcoding Policies: Never hardcode passwords, API keys, database URLs, or security secrets. Permissible hardcoded items are limited to non-sensitive defaults, such as default pagination size (10 items), UI label constants, and fixed navigation routes."),
        ("body", "Security Practices: .env and .env.local are permanently registered in .gitignore. Production secrets are restricted to designated system administrators. Secrets are rotated every 90 days or immediately upon team member offboarding."),

        # 4.2.5 Folder Structure
        ("heading", "4.2.5. Folder Structure"),
        ("body", "Organizing Principle: The repository employs a Layered, Decoupled Client-Server Modular Architecture. The codebase cleanly separates the user interface Single Page Application (/Frontend) from the stateless REST API (/Backend) and shared data migration assets (/database)."),
        ("body", "Top-Level Structure and Sub-Folder Responsibilities:"),
        ("bullet_bold", "/Backend/api/: ", "PHP RESTful API endpoints (donations, inventory, distributions, auth), session handling, and CORS middleware."),
        ("bullet_bold", "/Backend/api/uploads/: ", "Secure file system directory for recipient proof photos, donation receipts, and volunteer CV uploads."),
        ("bullet_bold", "/Frontend/src/pages/: ", "Top-level page views (Home, AboutUs, FAQ, Public Donation Tracking, Login, Register)."),
        ("bullet_bold", "/Frontend/src/portals/: ", "Role-based user portal layouts and specialized sub-views (Admin, Donor, Volunteer, Beneficiary, Staff)."),
        ("bullet_bold", "/Frontend/src/components/: ", "Reusable modular interface components divided into /admin, /layout, /shared, and /home."),
        ("bullet_bold", "/Frontend/src/context/: ", "Global React Context providers (AuthContext.jsx, ThemeContext.jsx)."),
        ("bullet_bold", "/Frontend/src/hooks/: ", "Reusable custom React hooks (useApiList.js, usePagination.js)."),
        ("bullet_bold", "/Frontend/src/styles/: ", "Global CSS stylesheets, CSS custom property design tokens, and role-specific portal styles."),
        ("bullet_bold", "/database/: ", "Master SQL schema files (donation_system.sql), seed fixtures, and DATABASE_METADATA.md documentation."),
        ("bullet_bold", "/mail-service/: ", "Express / Node.js background mail notification microservice for queueing transactional notices."),
        ("body", "Nested Rules and Naming: Components are grouped by feature domain. Any UI element reused across two or more portals must reside under /components/shared/. Folder names are lowercase with hyphens for services and PascalCase for React component folders."),

        # 4.2.6 Taxonomy/Metadata Standards
        ("heading", "4.2.6. Taxonomy/Metadata Standards"),
        ("body", "Purpose: Establishing rigorous taxonomy and metadata standards guarantees consistent data classification, predictable API contracts, structured code documentation, and unambiguous semantic versioning across the entire humanitarian relief system lifecycle."),
        ("bullet_bold", "Database Schema Taxonomy: ", "Tables are pluralized snake_case nouns (e.g., assistance_requests). Primary keys are consistently named id (BIGINT UNSIGNED AUTO_INCREMENT). Foreign keys strictly follow singular_parent_table_id. Audit timestamps created_at and updated_at are mandatory on all core transactional tables."),
        ("bullet_bold", "API Metadata and Envelope: ", "Every API endpoint returns explicit metadata headers (Content-Type: application/json; charset=utf-8, X-API-Version: 1.0.0). Paginated endpoints return metadata envelopes containing total_records, current_page, per_page, and total_pages."),
        ("bullet_bold", "Code Documentation Metadata: ", "Standardized PHPDoc and JSDoc annotations (@param, @returns, @throws, @deprecated, @author) must precede all core utility routines and custom hooks."),
        ("bullet_bold", "Version Control Tagging: ", "Semantic Versioning (SemVer 2.0.0) format (vMAJOR.MINOR.PATCH) is enforced (e.g., v1.0.0 Initial Capstone Release). Git commit messages follow Conventional Commits (feat:, fix:, docs:, refactor:, chore:)."),
        ("bullet_bold", "Feature and Environment Flags: ", "Runtime flags (e.g., FEATURE_AI_VOLUNTEER_MATCHING=true, MAINTENANCE_MODE=false) dictate module behavior without code modifications."),

        # 4.2.7 Deployment Standards
        ("heading", "4.2.7. Deployment Standards"),
        ("body", "Deployment Environments: 1) Development (Local machine localhost:5173 / localhost:8000) for rapid feature iteration; 2) UAT (Staging on Vercel and Render) for validation by Rise Above Foundation Cebu staff and SWU faculty panelists; and 3) Production (Vercel Global CDN + Render Docker + TiDB Cloud Serverless) for live humanitarian operations."),
        ("body", "CI/CD Deployment Process and Pipeline:"),
        ("bullet_bold", "1. Code Commit: ", "Developer pushes feature commits to GitHub repository branches."),
        ("bullet_bold", "2. Automated Trigger: ", "GitHub webhook notifies Render and Vercel build pipelines."),
        ("bullet_bold", "3. Build and Test: ", "Vercel executes 'npm run build' (Vite bundle optimization and lint validation); Render builds Docker container with PHP 8.2 and Apache."),
        ("bullet_bold", "4. Health Check Verification: ", "Render verifies zero-downtime container readiness at /api/health.php before shifting live traffic."),
        ("bullet_bold", "5. Database Migrations: ", "Schema modifications are applied using transaction-safe migration scripts against TiDB Cloud."),
        ("body", "Branching Strategy: A structured GitFlow model is implemented:"),
        ("bullet_bold", "feature/*: ", "Short-lived feature branches branched from develop."),
        ("bullet_bold", "develop: ", "Integration branch continuously deployed to the UAT / Staging environment."),
        ("bullet_bold", "main: ", "Production-ready protected branch. Pushing or merging directly to main is restricted."),
        ("bullet_bold", "hotfix/*: ", "Emergency patches branched from main and back-merged to develop."),
        ("body", "Approval and Review Requirements: Merging into main requires a formal Pull Request (PR) with at least two peer reviews from capstone project members and validation that automated build checks succeed. Final approval is signed off by the Lead Developer."),
        ("body", "Rollback Procedures:"),
        ("bullet_bold", "Frontend (Vercel): ", "Access Vercel Dashboard -> Deployments -> locate previous verified deployment -> click 'Rollback to this Deployment'. Rollback executes instantaneously within 15 seconds."),
        ("bullet_bold", "Backend (Render): ", "Access Render Dashboard -> Service Activity -> select prior working Docker image build -> trigger redeploy."),
        ("bullet_bold", "Database (TiDB Cloud): ", "Point-In-Time Recovery (PITR) snapshots allow restoring the database state to any exact second prior to the failed deployment.")
    ]

    p_start_idx = None
    p_end_idx = None
    for idx, p in enumerate(doc.paragraphs):
        if "4.2.2. Code Naming Convention" in p.text:
            p_start_idx = idx
        if "4.2.8. Permissions & Access Management" in p.text:
            p_end_idx = idx
            break
            
    paras_to_remove = doc.paragraphs[p_start_idx : p_end_idx]
    end_elem = doc.paragraphs[p_end_idx]._p
    parent = end_elem.getparent()
    
    for item_type, *item_args in gov_content:
        new_p = doc.add_paragraph()
        if item_type == "heading":
            r = new_p.add_run(item_args[0])
            format_run(r, bold=True, size=11)
            new_p.paragraph_format.space_before = Pt(12)
            new_p.paragraph_format.space_after = Pt(4)
        elif item_type == "body":
            r = new_p.add_run(item_args[0])
            format_run(r, bold=False, size=11)
            new_p.paragraph_format.space_after = Pt(6)
            new_p.paragraph_format.line_spacing = 1.15
        elif item_type == "bullet_bold":
            new_p.paragraph_format.left_indent = Inches(0.25)
            r_bullet = new_p.add_run("• ")
            format_run(r_bullet, bold=False, size=11)
            
            r_bold = new_p.add_run(item_args[0])
            format_run(r_bold, bold=True, size=11)
            
            r_text = new_p.add_run(item_args[1])
            format_run(r_text, bold=False, size=11)
            
            new_p.paragraph_format.space_after = Pt(3)
            new_p.paragraph_format.line_spacing = 1.15
            
        parent.insert(parent.index(end_elem), new_p._p)

    for p in paras_to_remove:
        p_elem = p._p
        if p_elem.getparent() is not None:
            p_elem.getparent().remove(p_elem)

    # Save to Downloads
    doc.save(src_path)
    print(f"Updated in Downloads: {src_path}")

    # Save to Desktop
    desktop_copy = r"C:\Users\James\OneDrive\Desktop\SWU 003 (Final Project) System Design & Governance Document.docx"
    doc.save(desktop_copy)
    print(f"Updated on Desktop: {desktop_copy}")

    # Save to Project Workspace
    ws_copy = r"c:\Users\James\DonationSystem\SWU 003 (Final Project) System Design & Governance Document.docx"
    doc.save(ws_copy)
    print(f"Updated in Workspace: {ws_copy}")

update_manuscript()
